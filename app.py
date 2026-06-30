from __future__ import annotations

import os
import threading
import time
import unicodedata
from datetime import datetime
from io import BytesIO
from typing import Any

import pandas as pd
import requests
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

API_BASE = "https://opendata-api.businessportal.gr/api/opendata/v1"
REQUEST_INTERVAL_SECONDS = 8.2  # Κάτω από το όριο 8 αιτημάτων/λεπτό.
PAGE_SIZE = 200
MAX_ACTIVITY_IDS = 100

st.set_page_config(
    page_title="Αναζήτηση επιχειρήσεων ΓΕΜΗ",
    page_icon="🔎",
    layout="wide",
)


# -----------------------------
# Βοηθητικές συναρτήσεις
# -----------------------------
def normalize_text(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return " ".join(text.upper().split())


def kad_digits(value: Any) -> str:
    return "".join(char for char in str(value or "") if char.isdigit())


def first(obj: Any, *keys: str, default: Any = "") -> Any:
    if not isinstance(obj, dict):
        return default
    for key in keys:
        value = obj.get(key)
        if value not in (None, "", [], {}):
            return value
    return default


def description(value: Any) -> str:
    if isinstance(value, dict):
        return str(first(value, "descr", "description", "name", "title"))
    return str(value or "")


def as_list(payload: Any, *keys: str) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        return [item for item in payload if isinstance(item, dict)]
    if isinstance(payload, dict):
        for key in (*keys, "searchResults", "results", "items", "data"):
            value = payload.get(key)
            if isinstance(value, list):
                return [item for item in value if isinstance(item, dict)]
    return []


def parse_kad_prefixes(value: str) -> list[str]:
    raw_parts = value.replace(";", ",").split(",")
    prefixes: list[str] = []
    for part in raw_parts:
        prefix = kad_digits(part)
        if prefix and prefix not in prefixes:
            prefixes.append(prefix)
    return prefixes


def profession_stems(value: str) -> list[str]:
    """Δημιουργεί απλές ρίζες ώστε π.χ. ΥΔΡΑΥΛΙΚΟΣ να βρίσκει ΥΔΡΑΥΛΙΚΕΣ."""
    ignored = {"ΚΑΙ", "ΓΙΑ", "ΤΩΝ", "ΤΗΣ", "ΤΟ", "Η", "Ο", "ΣΕ", "ΜΕ"}
    stems: list[str] = []
    for token in normalize_text(value).split():
        if token in ignored or len(token) < 4:
            continue
        if len(token) >= 10:
            stem = token[:7]
        elif len(token) >= 7:
            stem = token[:6]
        else:
            stem = token
        if stem not in stems:
            stems.append(stem)
    return stems


def get_api_key() -> str:
    # 1) Streamlit Secrets, 2) μεταβλητή περιβάλλοντος, 3) προσωρινή εισαγωγή χρήστη.
    try:
        secret_key = str(st.secrets.get("GEMI_API_KEY", "")).strip()
    except Exception:
        secret_key = ""

    env_key = os.getenv("GEMI_API_KEY", "").strip()
    return secret_key or env_key


class GemiClient:
    def __init__(self, api_key: str):
        self.session = requests.Session()
        self.session.headers.update({
            "api_key": api_key,
            "Accept": "application/json",
        })
        self.lock = threading.Lock()
        self.last_call = 0.0

    def get(self, path: str, params: dict[str, Any] | None = None, allow_404: bool = False) -> Any:
        url = API_BASE + path

        for attempt in range(4):
            try:
                with self.lock:
                    elapsed = time.monotonic() - self.last_call
                    if self.last_call and elapsed < REQUEST_INTERVAL_SECONDS:
                        time.sleep(REQUEST_INTERVAL_SECONDS - elapsed)

                    response = self.session.get(url, params=params, timeout=90)
                    self.last_call = time.monotonic()

                if response.status_code == 401:
                    raise RuntimeError("401 Unauthorized: το API key δεν έγινε αποδεκτό.")

                if response.status_code == 404 and allow_404:
                    return {}

                if response.status_code == 429:
                    if attempt == 3:
                        raise RuntimeError("429 Too Many Requests: εξαντλήθηκε το όριο κλήσεων.")
                    time.sleep(65)
                    continue

                if response.status_code >= 500:
                    if attempt == 3:
                        raise RuntimeError(
                            f"Προσωρινό σφάλμα ΓΕΜΗ ({response.status_code})."
                        )
                    time.sleep(20)
                    continue

                if response.status_code >= 400:
                    raise RuntimeError(
                        f"Σφάλμα API {response.status_code}: {response.text[:500]}"
                    )

                return response.json()

            except requests.RequestException as exc:
                if attempt == 3:
                    raise RuntimeError(f"Αποτυχία σύνδεσης με το ΓΕΜΗ: {exc}") from exc
                time.sleep(15)

        raise RuntimeError("Δεν ολοκληρώθηκε η κλήση στο API.")


@st.cache_resource(show_spinner=False)
def get_client(api_key: str) -> GemiClient:
    return GemiClient(api_key)


@st.cache_data(ttl=24 * 60 * 60, show_spinner=False)
def load_metadata(_client: GemiClient, path: str, list_key: str) -> list[dict[str, Any]]:
    return as_list(_client.get(path), list_key)


def find_activities(
    activities: list[dict[str, Any]],
    profession: str,
    kad_prefixes: list[str],
) -> list[dict[str, Any]]:
    stems = profession_stems(profession)
    matches: list[dict[str, Any]] = []

    for item in activities:
        code = kad_digits(first(item, "id", "code", "kad"))
        text = normalize_text(first(item, "descr", "description", "name"))

        # Αν δόθηκε ΚΑΔ, αυτός έχει προτεραιότητα και η ειδικότητα λειτουργεί ως ετικέτα.
        if kad_prefixes:
            matched = any(code.startswith(prefix) for prefix in kad_prefixes)
        else:
            matched = bool(stems) and all(stem in text for stem in stems)

        if matched:
            matches.append(item)

    unique: dict[str, dict[str, Any]] = {}
    for item in matches:
        item_id = str(first(item, "id", "code", "kad")).strip()
        if item_id:
            unique[item_id] = item

    ordered = sorted(
        unique.values(),
        key=lambda item: (
            len(kad_digits(first(item, "id", "code", "kad"))),
            str(first(item, "id", "code", "kad")),
        ),
    )
    return ordered[:MAX_ACTIVITY_IDS]


def find_location(
    locations: list[dict[str, Any]],
    area_text: str,
) -> tuple[dict[str, Any] | None, list[dict[str, Any]]]:
    target = normalize_text(area_text)
    if not target:
        return None, []

    exact = [
        item for item in locations
        if normalize_text(first(item, "descr", "description", "name")) == target
    ]
    contains = [
        item for item in locations
        if target in normalize_text(first(item, "descr", "description", "name"))
    ]

    candidates = exact or contains
    candidates.sort(
        key=lambda item: len(normalize_text(first(item, "descr", "description", "name")))
    )
    return (candidates[0] if candidates else None), candidates


def company_key(company: dict[str, Any]) -> str:
    return str(first(
        company,
        "arGemi", "gemiNumber", "gemi_number",
        "afm", "vatNumber",
        "coNameEl", "coName", "companyName", "name",
    ))


def search_companies(
    client: GemiClient,
    activity_ids: list[str],
    location_parameter: str,
    location_id: str,
    only_active: bool,
    max_results: int,
) -> list[dict[str, Any]]:
    companies: dict[str, dict[str, Any]] = {}

    # Μικρές ομάδες ΚΑΔ για να μην γίνει υπερβολικά μεγάλο query string.
    for start in range(0, len(activity_ids), 20):
        group = activity_ids[start:start + 20]
        offset = 0

        while len(companies) < max_results:
            params = {
                "activities": ",".join(group),
                location_parameter: location_id,
                "isActive": str(only_active).lower(),
                "resultsOffset": offset,
                "resultsSize": min(PAGE_SIZE, max_results - len(companies)),
                "resultsSortBy": "+coName",
            }

            payload = client.get("/companies", params=params, allow_404=True)
            page = as_list(payload, "searchResults", "companies")
            if not page:
                break

            for company in page:
                key = company_key(company)
                if key:
                    companies[key] = company
                if len(companies) >= max_results:
                    break

            if len(page) < params["resultsSize"] or len(companies) >= max_results:
                break
            offset += len(page)

        if len(companies) >= max_results:
            break

    return list(companies.values())


def company_activity_text(company: dict[str, Any]) -> tuple[str, str]:
    codes: list[str] = []
    texts: list[str] = []
    values = company.get("activities") or []
    if isinstance(values, dict):
        values = [values]

    for item in values if isinstance(values, list) else []:
        if not isinstance(item, dict):
            continue
        value = item.get("activity") if isinstance(item.get("activity"), dict) else item
        code = first(value, "id", "code", "kad")
        text = first(value, "descr", "description", "name")
        if code:
            codes.append(str(code))
        if text:
            texts.append(str(text))

    return " | ".join(dict.fromkeys(codes)), " | ".join(dict.fromkeys(texts))


def companies_to_dataframe(companies: list[dict[str, Any]]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []

    for company in companies:
        codes, activity_descriptions = company_activity_text(company)
        street = str(first(company, "street", "streetName")).strip()
        number = str(first(company, "streetNumber", "number")).strip()
        address = " ".join(part for part in [street, number] if part)

        rows.append({
            "Επωνυμία": first(company, "coNameEl", "coName", "companyName", "name"),
            "Διακριτικός τίτλος": " | ".join(
                str(value) for value in (company.get("coTitlesEl") or [])
            ) if isinstance(company.get("coTitlesEl"), list) else first(company, "coTitlesEl", "tradeName"),
            "Αριθμός ΓΕΜΗ": first(company, "arGemi", "gemiNumber", "gemi_number"),
            "ΑΦΜ": first(company, "afm", "vatNumber", "vat"),
            "Κατάσταση": description(first(company, "status")),
            "Νομική μορφή": description(first(company, "legalType", "legalForm")),
            "Δήμος": description(first(company, "municipality")),
            "Νομός": description(first(company, "prefecture")),
            "Πόλη": first(company, "city"),
            "Διεύθυνση": address,
            "ΤΚ": first(company, "zipCode", "postalCode"),
            "Email": first(company, "email"),
            "Website": first(company, "url", "website"),
            "Ημερομηνία ίδρυσης": first(company, "incorporationDate"),
            "ΚΑΔ": codes,
            "Περιγραφές ΚΑΔ": activity_descriptions,
            "Υποκατάστημα": "Ναι" if company.get("isBranch") else "Όχι",
        })

    columns = [
        "Επωνυμία", "Διακριτικός τίτλος", "Αριθμός ΓΕΜΗ", "ΑΦΜ",
        "Κατάσταση", "Νομική μορφή", "Δήμος", "Νομός", "Πόλη",
        "Διεύθυνση", "ΤΚ", "Email", "Website", "Ημερομηνία ίδρυσης",
        "ΚΑΔ", "Περιγραφές ΚΑΔ", "Υποκατάστημα",
    ]
    dataframe = pd.DataFrame(rows, columns=columns)
    if not dataframe.empty:
        dataframe = dataframe.sort_values(
            by=["Επωνυμία", "Αριθμός ΓΕΜΗ"],
            na_position="last",
        ).reset_index(drop=True)
    return dataframe


def create_excel(
    dataframe: pd.DataFrame,
    profession: str,
    area_name: str,
    area_type: str,
    requested_kad: str,
    matched_activities: list[dict[str, Any]],
) -> bytes:
    output = BytesIO()
    kad_lines = "\n".join(
        f"{first(item, 'id', 'code', 'kad')} — "
        f"{first(item, 'descr', 'description', 'name')}"
        for item in matched_activities
    )

    info = pd.DataFrame({
        "Πεδίο": [
            "Ειδικότητα", "Περιοχή", "Τύπος περιοχής", "ΚΑΔ που πληκτρολογήθηκε",
            "Πλήθος αποτελεσμάτων", "ΚΑΔ αναζήτησης", "Ημερομηνία δημιουργίας",
            "Πηγή", "Άδεια δεδομένων",
        ],
        "Τιμή": [
            profession or "—", area_name, area_type, requested_kad or "Αυτόματη εύρεση",
            len(dataframe), kad_lines, datetime.now().strftime("%d/%m/%Y %H:%M"),
            "OpenData ΓΕΜΗ — https://opendata.businessportal.gr/", "ODC-BY-1.0",
        ],
    })

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        dataframe.to_excel(writer, sheet_name="Επιχειρήσεις", index=False)
        info.to_excel(writer, sheet_name="Πληροφορίες", index=False)

        workbook = writer.book
        header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)

        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for column_cells in worksheet.columns:
                letter = get_column_letter(column_cells[0].column)
                max_length = 0
                for cell in column_cells:
                    value = "" if cell.value is None else str(cell.value)
                    max_length = max(
                        max_length,
                        max((len(line) for line in value.split("\n")), default=0),
                    )
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                worksheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 45)

        results_sheet = workbook["Επιχειρήσεις"]
        for row_number in range(2, results_sheet.max_row + 1):
            for column_number in (3, 4, 11):
                results_sheet.cell(row=row_number, column=column_number).number_format = "@"

    output.seek(0)
    return output.getvalue()


# -----------------------------
# Περιβάλλον ιστοσελίδας
# -----------------------------
st.title("🔎 Αναζήτηση επιχειρήσεων στο ΓΕΜΗ")
st.caption("Αναζήτηση με ειδικότητα ή ΚΑΔ και γεωγραφική περιοχή, με εξαγωγή σε Excel.")

configured_key = get_api_key()
with st.sidebar:
    st.header("Ρυθμίσεις API")
    if configured_key:
        st.success("Το API key έχει ρυθμιστεί με ασφάλεια στον server.")
        api_key = configured_key
    else:
        api_key = st.text_input(
            "API key ΓΕΜΗ",
            type="password",
            help="Το κλειδί χρησιμοποιείται μόνο για την τρέχουσα συνεδρία.",
        ).strip()
        st.warning("Για δημόσια εγκατάσταση βάλε το key στα Streamlit Secrets.")

    st.divider()
    st.caption("Πηγή: OpenData ΓΕΜΗ · Άδεια ODC-BY-1.0")

with st.form("search_form"):
    col1, col2, col3 = st.columns(3)

    with col1:
        profession = st.text_input(
            "Ειδικότητα / κλάδος",
            value="Υδραυλικός",
            placeholder="π.χ. Υδραυλικός",
            help="Χρησιμοποιείται για αυτόματη εύρεση ΚΑΔ όταν το πεδίο ΚΑΔ είναι κενό.",
        ).strip()

    with col2:
        area = st.text_input(
            "Περιοχή",
            value="Θεσσαλονίκη",
            placeholder="π.χ. Θεσσαλονίκη",
        ).strip()

    with col3:
        kad_input = st.text_input(
            "ΚΑΔ (προαιρετικός)",
            value="43.22",
            placeholder="π.χ. 43.22 ή 43.22, 69.20",
            help="Μπορείς να βάλεις έναν ή περισσότερους ΚΑΔ. Ο ΚΑΔ έχει προτεραιότητα από την ειδικότητα.",
        ).strip()

    col4, col5, col6 = st.columns(3)
    with col4:
        area_mode_label = st.selectbox(
            "Τύπος περιοχής",
            ["Δήμος", "Νομός / Περιφερειακή Ενότητα"],
        )
    with col5:
        only_active = st.checkbox("Μόνο ενεργές επιχειρήσεις", value=True)
    with col6:
        max_results = st.selectbox("Μέγιστος αριθμός αποτελεσμάτων", [200, 400, 600, 1000])

    submitted = st.form_submit_button("Αναζήτηση στο ΓΕΜΗ", type="primary", use_container_width=True)

if submitted:
    if not api_key:
        st.error("Βάλε πρώτα το API key του ΓΕΜΗ.")
        st.stop()
    if not area:
        st.error("Συμπλήρωσε περιοχή.")
        st.stop()
    if not profession and not kad_input:
        st.error("Συμπλήρωσε ειδικότητα ή ΚΑΔ.")
        st.stop()

    client = get_client(api_key)
    location_path = "/metadata/municipalities" if area_mode_label == "Δήμος" else "/metadata/prefectures"
    location_list_key = "municipalities" if area_mode_label == "Δήμος" else "prefectures"
    location_parameter = "municipalities" if area_mode_label == "Δήμος" else "prefectures"

    try:
        with st.status("Επικοινωνία με το ΓΕΜΗ…", expanded=True) as status:
            st.write("Λήψη καταλόγου ΚΑΔ…")
            activities = load_metadata(client, "/metadata/activities", "activities")

            st.write("Εύρεση περιοχής…")
            locations = load_metadata(client, location_path, location_list_key)
            selected_location, location_candidates = find_location(locations, area)
            if not selected_location:
                raise RuntimeError(f"Δεν βρέθηκε περιοχή που να ταιριάζει με «{area}».")

            location_id = str(first(selected_location, "id", "code"))
            location_name = str(first(selected_location, "descr", "description", "name"))
            st.write(f"Επιλέχθηκε: **{location_name}**")

            kad_prefixes = parse_kad_prefixes(kad_input)
            matched_activities = find_activities(activities, profession, kad_prefixes)
            if not matched_activities:
                raise RuntimeError(
                    "Δεν βρέθηκαν σχετικοί ΚΑΔ. Δοκίμασε γενικότερη ειδικότητα ή έλεγξε τον ΚΑΔ."
                )

            activity_ids = [
                str(first(item, "id", "code", "kad"))
                for item in matched_activities
            ]
            st.write(f"Βρέθηκαν **{len(activity_ids)}** σχετικοί ΚΑΔ.")

            st.write("Αναζήτηση επιχειρήσεων…")
            companies = search_companies(
                client=client,
                activity_ids=activity_ids,
                location_parameter=location_parameter,
                location_id=location_id,
                only_active=only_active,
                max_results=int(max_results),
            )

            dataframe = companies_to_dataframe(companies)
            excel_bytes = create_excel(
                dataframe=dataframe,
                profession=profession,
                area_name=location_name,
                area_type=area_mode_label,
                requested_kad=kad_input,
                matched_activities=matched_activities,
            )
            status.update(label="Η αναζήτηση ολοκληρώθηκε.", state="complete", expanded=False)

        st.success(f"Βρέθηκαν {len(dataframe)} μοναδικές επιχειρήσεις.")

        if len(location_candidates) > 1:
            alternatives = ", ".join(
                str(first(item, "descr", "description", "name"))
                for item in location_candidates[1:6]
            )
            if alternatives:
                st.info(f"Άλλες πιθανές αντιστοιχίσεις περιοχής: {alternatives}")

        with st.expander("ΚΑΔ που χρησιμοποιήθηκαν"):
            for item in matched_activities:
                st.write(
                    f"**{first(item, 'id', 'code', 'kad')}** — "
                    f"{first(item, 'descr', 'description', 'name')}"
                )

        st.dataframe(dataframe, use_container_width=True, hide_index=True)
        filename_area = normalize_text(area).replace(" ", "_").lower() or "perioxi"
        st.download_button(
            "⬇️ Λήψη αποτελεσμάτων σε Excel",
            data=excel_bytes,
            file_name=f"gemi_{filename_area}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
            use_container_width=True,
        )

    except Exception as exc:
        st.error(str(exc))
        st.info(
            "Έλεγξε το API key, τα κριτήρια αναζήτησης και αν έχει ξεπεραστεί το όριο "
            "των 8 αιτημάτων ανά λεπτό."
        )
